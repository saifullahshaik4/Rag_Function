import os
import logging
import io
from typing import List, Dict
from fastapi import FastAPI, UploadFile, File, HTTPException, Body
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from .processing import init_gemini, process_single_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from dotenv import load_dotenv
load_dotenv()

# Setup logger for main API
logger = logging.getLogger(__name__)

app = FastAPI(title="RAG User Story API", version="0.1")

# CORS for local static frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5500", "http://127.0.0.1:5500", "http://localhost:5173", "http://localhost:3000", "*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def _startup():
    logger.info("Application startup initiated")
    try:
        init_gemini()
        logger.info("Application startup completed successfully")
    except Exception as e:
        logger.error(f"Application startup failed: {str(e)}", exc_info=True)
        raise

@app.get("/api/health")
def health():
    logger.debug("Health check endpoint called")
    return {"ok": True}

@app.post("/api/generate")
async def generate(file: UploadFile = File(...), max_stories: int = 3):
    logger.info(f"Generate endpoint called - File: {file.filename}, Max stories: {max_stories}")
    
    # Check file size
    if file.size and file.size > 10 * 1024 * 1024:
        logger.warning(f"File rejected - too large: {file.size} bytes ({file.filename})")
        raise HTTPException(status_code=413, detail="File too large (>10MB)")
    
    logger.info(f"Reading file: {file.filename} ({file.size if file.size else 'unknown'} bytes)")

    content = await file.read()
    if not content:
        logger.error(f"Empty file uploaded: {file.filename}")
        raise HTTPException(status_code=400, detail="Empty file")

    try:
        logger.info(f"Starting processing of {file.filename}")
        result = process_single_file(content, file.filename, max_stories=max_stories)
        
        # Log summary of results
        num_stories = len(result.get("stories", []))
        num_errors = sum(1 for s in result.get("stories", []) if "error" in s)
        logger.info(f"Processing completed - Generated {num_stories} stories ({num_errors} errors) for {file.filename}")
        
        return JSONResponse(result)
    except Exception as e:
        logger.error(f"Fatal error processing {file.filename}: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/export/xlsx")
async def export_xlsx(stories: List[Dict] = Body(...)):
    """
    Export user stories to XLSX format with columns:
    Story Number, Summary, Description, Acceptance Criteria, Non Functional, Priority, Tag
    """
    logger.info(f"Export XLSX endpoint called with {len(stories)} stories")
    
    try:
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "User Stories"
        
        # Define column headers
        headers = [
            "Story Number",
            "Summary", 
            "Description",
            "Acceptance Criteria",
            "Non Functional",
            "Priority",
            "Tag"
        ]
        
        # Style the header row
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Set column widths
        column_widths = {
            'A': 15,  # Story Number
            'B': 35,  # Summary
            'C': 50,  # Description
            'D': 50,  # Acceptance Criteria
            'E': 30,  # Non Functional
            'F': 15,  # Priority
            'G': 25,  # Tag
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # Content alignment
        content_alignment = Alignment(vertical="top", wrap_text=True)
        
        # Write story data
        story_num = 0
        for story in stories:
            # Skip error stories
            if "error" in story:
                continue
            
            story_num += 1
            row_num = story_num + 1  # +1 for header row
            
            # Story Number
            ws.cell(row=row_num, column=1, value=story_num)
            
            # Summary (title)
            ws.cell(row=row_num, column=2, value=story.get("title", ""))
            
            # Description (story)
            ws.cell(row=row_num, column=3, value=story.get("story", ""))
            
            # Acceptance Criteria (join list with newlines)
            ac_list = story.get("acceptance_criteria", [])
            ac_text = "\n".join(f"- {ac}" for ac in ac_list) if ac_list else ""
            ws.cell(row=row_num, column=4, value=ac_text)
            
            # Non Functional (join list with commas)
            nf_list = story.get("non_functional", [])
            nf_text = ", ".join(nf_list) if nf_list else ""
            ws.cell(row=row_num, column=5, value=nf_text)
            
            # Priority
            ws.cell(row=row_num, column=6, value=story.get("priority", ""))
            
            # Tags (join list with commas)
            tags_list = story.get("tags", [])
            tags_text = ", ".join(tags_list) if tags_list else ""
            ws.cell(row=row_num, column=7, value=tags_text)
            
            # Apply alignment to all cells in the row
            for col_num in range(1, 8):
                ws.cell(row=row_num, column=col_num).alignment = content_alignment
        
        # Freeze the header row
        ws.freeze_panes = "A2"
        
        # Save to bytes buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        logger.info(f"Successfully generated XLSX with {story_num} stories")
        
        # Return as streaming response
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=user_stories.xlsx"
            }
        )
    
    except Exception as e:
        logger.error(f"Error generating XLSX: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Error generating XLSX: {str(e)}")